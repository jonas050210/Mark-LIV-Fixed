"""Document handling: PDF, Word, plain text and Markdown, PowerPoint."""
from __future__ import annotations

from actions.file_handlers.archives import _archive_members
from actions.file_handlers.common import (
    Path,
    _bool_param,
    _file_size_str,
    _gemini_client,
    _output_path,
    _read_text_prefix,
    _source_stem,
    _write_generated,
    atomic_create_text,
    islice,
)


def _process_pdf(path: Path, action: str, params: dict, speak=None) -> str:
    action = action or "summarize"

    def _extract_pdf_text(max_chars=50_000, max_pages=200) -> str:
        chunks = []
        length = 0
        try:
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                pages = pdf.pages[:max_pages]
                for page in pages:
                    chunk = (page.extract_text() or "") + "\n"
                    chunks.append(chunk)
                    length += len(chunk)
                    if length >= max_chars:
                        break
        except ImportError:
            try:
                import pypdf
                with open(path, "rb") as handle:
                    reader = pypdf.PdfReader(handle)
                    for page in islice(reader.pages, max_pages):
                        chunk = (page.extract_text() or "") + "\n"
                        chunks.append(chunk)
                        length += len(chunk)
                        if length >= max_chars:
                            break
            except ImportError:
                return ""
        return "".join(chunks)[:max_chars]

    if action in ("summarize", "extract_text", "translate_hint", "analyze", "reformat"):
        text = _extract_pdf_text()
        if not text.strip():
            return "Could not extract text from PDF (may be scanned/image-based)."

        if action == "extract_text":
            out = _output_path(path, "text", ".txt")
            atomic_create_text(out, text)
            return f"Text extracted ({len(text)} chars). Saved: {out.name}"

        prompt_map = {
            "summarize":      f"Summarize this PDF document concisely:\n\n{text}",
            "analyze":        f"Analyze this document thoroughly:\n\n{text}",
            "translate_hint": f"What language is this document in and what does it say? Summarize:\n\n{text}",
            "reformat":       f"Reformat this text cleanly with proper structure:\n\n{text}",
        }
        try:
            model    = _gemini_client()
            response = model.generate_content(prompt_map.get(action, f"Analyze:\n\n{text}"))
            result   = response.text.strip()
            if len(result) > 600 and _bool_param(params, "save", True):
                out = _output_path(path, action, ".txt")
                atomic_create_text(out, result)
                return f"{result[:400]}...\n\nFull result saved: {out.name}"
            return result
        except Exception as e:
            return f"AI analysis failed: {type(e).__name__}"

    if action == "info":
        try:
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                pages = len(pdf.pages)
            return f"PDF: {pages} pages, size: {_file_size_str(path)}"
        except Exception:
            return f"PDF size: {_file_size_str(path)}"

    if action == "to_word":
        text = _extract_pdf_text()
        if not text:
            return "Could not extract text to convert."
        try:
            from docx import Document
            doc  = Document()
            doc.add_heading(_source_stem(path), 0)
            for para in text.split("\n\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())
            out = _output_path(path, "converted", ".docx")
            _write_generated(out, doc.save)
            return f"Converted to Word document. Saved: {out.name}"
        except ImportError:
            return "python-docx not installed. Run: pip install python-docx"

    return f"Unknown PDF action: '{action}'. Try: summarize, extract_text, info, to_word"
def _process_text_doc(path: Path, file_type: str, action: str,
                       params: dict, speak=None) -> str:
    action = action or "summarize"

    def _read_content() -> str:
        if file_type == "docx":
            try:
                if path.suffix.lower() == ".docx":
                    _archive_members(path)
                from docx import Document
                doc = Document(path)
                chunks = []
                length = 0
                for paragraph in doc.paragraphs:
                    chunk = str(paragraph.text)[:10_000]
                    chunks.append(chunk)
                    length += len(chunk) + 1
                    if length >= 50_000:
                        break
                return "\n".join(chunks)[:50_000]
            except ImportError:
                return "python-docx not installed."
            except ValueError as e:
                return f"Office document rejected: {type(e).__name__}"
            except Exception as e:
                return f"Read failed: {type(e).__name__}"
        else:
            return _read_text_prefix(path, max_bytes=2_000_000)

    content = _read_content()
    if content.startswith(("Office document rejected:", "Read failed:", "python-docx not installed.")):
        return content
    if not content.strip():
        return "File appears to be empty."

    if action == "word_count":
        words = len(content.split())
        chars = len(content)
        lines = content.count("\n")
        return f"Word count: {words} words, {chars} characters, {lines} lines."

    if action == "extract_text":
        if file_type != "txt":
            out = _output_path(path, "extracted", ".txt")
            atomic_create_text(out, content)
            return f"Text extracted. Saved: {out.name}"
        return content[:2000]

    instruction = params.get("instruction", "")
    prompt_map  = {
        "summarize":  f"Summarize this document concisely:\n\n{content[:40000]}",
        "analyze":    f"Analyze this document:\n\n{content[:40000]}",
        "reformat":   f"Reformat this text with clean structure, proper headings and paragraphs:\n\n{content[:40000]}",
        "fix":        f"Fix grammar, spelling and style issues in this text:\n\n{content[:40000]}",
        "translate_hint": f"What language is this and what does it say? Summarize:\n\n{content[:10000]}",
        "to_bullet":  f"Convert this text into a clear bullet-point summary:\n\n{content[:40000]}",
        "custom":     f"{instruction}\n\n{content[:40000]}",
    }

    if action not in prompt_map:

        action  = "custom"
        instruction = action

    try:
        model    = _gemini_client()
        response = model.generate_content(prompt_map[action])
        result   = response.text.strip()
        if len(result) > 600 and _bool_param(params, "save", True):
            out = _output_path(path, action, ".txt")
            atomic_create_text(out, result)
            return f"{result[:400]}...\n\nFull result saved: {out.name}"
        return result
    except Exception as e:
        return f"AI processing failed: {type(e).__name__}"
def _process_pptx(path: Path, action: str, params: dict, speak=None) -> str:
    action = action or "summarize"

    def _read_pptx_text() -> str:
        try:
            if path.suffix.lower() == ".pptx":
                _archive_members(path)
            from pptx import Presentation
            prs = Presentation(path)
            text = []
            total = 0
            for i, slide in enumerate(prs.slides, 1):
                if i > 200 or total >= 50_000:
                    break
                slide_parts = [f"\n--- Slide {i} ---\n"]
                for shape in slide.shapes:
                    value = getattr(shape, "text", "")
                    if isinstance(value, str) and value.strip():
                        slide_parts.append(value.strip()[:10_000] + "\n")
                slide_text = "".join(slide_parts)
                text.append(slide_text)
                total += len(slide_text)
            return "\n".join(text)[:50_000]
        except ImportError:
            return "python-pptx not installed."
        except ValueError as exc:
            return f"Office document rejected: {type(exc).__name__}"
        except Exception as exc:
            return f"Presentation read failed: {type(exc).__name__}"

    if action in ("summarize", "extract_text", "analyze"):
        text = _read_pptx_text()
        if text.startswith(("Office document rejected:", "Presentation read failed:", "python-pptx not installed.")):
            return text
        if action == "extract_text":
            out = _output_path(path, "text", ".txt")
            atomic_create_text(out, text)
            return f"Text extracted. Saved: {out.name}"
        try:
            model    = _gemini_client()
            prompt   = f"{'Summarize' if action == 'summarize' else 'Analyze'} this presentation:\n{text[:30000]}"
            response = model.generate_content(prompt)
            return response.text.strip()
        except Exception as e:
            return f"AI processing failed: {type(e).__name__}"

    return f"Unknown PPTX action: '{action}'. Try: summarize, extract_text, analyze"
