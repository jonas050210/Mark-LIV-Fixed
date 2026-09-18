"""
document_qa.py — answer questions about a document without touching it.

Why this exists next to file_processor
──────────────────────────────────────
actions/file_processor.py converts, resizes, edits, splits and *writes* files,
and the action/plugin audit classifies it as `needs rewrite`. A large part of
what people actually ask — "what does this contract say about notice periods",
"which figure is on slide four" — needs no mutation at all. This action is the
strictly read-only half: it opens a document, extracts its text and answers.

It never writes, never creates output files, and never executes anything it
reads. Every file is opened read-only, and the path is resolved through
symlinks and must stay inside the user's own directories, so a spoken path
cannot be turned into a read of a system file.

Declares itself NON_BLOCKING: extracting a long PDF can take a while and the
model should not sit silent waiting for it. The answer re-enters the
conversation at the next gap in speech (WHEN_IDLE).
"""
from __future__ import annotations

import tempfile
from pathlib import Path

_MAX_CHARS = 120_000          # what we hand to the model
_MAX_FILE_BYTES = 40_000_000  # refuse before opening something enormous
_MAX_SHEET_ROWS = 400
_ANSWER_TIMEOUT_MS = 90_000

_SAFE_ROOTS = (Path.home(), Path(tempfile.gettempdir()))
_PLAIN_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv",
                   ".json", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".xml"}
_KNOWN_SUFFIXES = _PLAIN_SUFFIXES | {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx"}


def _safe_file(raw) -> Path | None:
    """Resolve a model-supplied path, refusing anything outside the user's space."""
    text = str(raw or "").strip().strip('"').strip("'")
    if not text:
        return None
    try:
        resolved = Path(text).expanduser().resolve()
    except Exception:
        return None
    for root in _SAFE_ROOTS:
        try:
            root_resolved = root.resolve()
            if resolved == root_resolved or resolved.is_relative_to(root_resolved):
                return resolved
        except Exception:
            continue
    return None


def _page_set(raw, total: int) -> set[int] | None:
    """Parse '3', '1-5' or '2,4,6-8' into 1-based page numbers.

    Returns None for "no restriction" and an empty set for "unparsable", so the
    caller can tell the two apart instead of silently reading the wrong pages.
    """
    text = str(raw or "").strip()
    if not text:
        return None
    wanted: set[int] = set()
    for part in text.replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            bounds = part.split("-", 1)
            try:
                lo, hi = int(bounds[0]), int(bounds[1])
            except ValueError:
                return set()
            if lo > hi:
                lo, hi = hi, lo
            wanted.update(range(lo, hi + 1))
        else:
            try:
                wanted.add(int(part))
            except ValueError:
                return set()
    return {n for n in wanted if 1 <= n <= max(total, 1)} or set()


def _extract_pdf(path: Path, pages_raw) -> tuple[str, str]:
    """PyMuPDF first (fastest C engine, best reading order), pdfplumber second, PyPDF2 fallback. Strictly read-only."""
    # Tier 1: PyMuPDF (pymupdf / fitz) - 10-50x faster and clean text extraction
    try:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
        with fitz.open(str(path)) as doc:
            total = len(doc)
            wanted = _page_set(pages_raw, total)
            if wanted is not None and not wanted:
                return "", "That page range is outside the document."
            chunks = []
            for number in range(1, total + 1):
                if wanted is not None and number not in wanted:
                    continue
                page = doc[number - 1]
                chunks.append(page.get_text() or "")
        text = "\n".join(chunks)
        if text.strip():
            return text, ""
    except ImportError:
        pass
    except Exception as exc:
        pass

    # Tier 2: pdfplumber (layout-aware fallback)
    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            total = len(pdf.pages)
            wanted = _page_set(pages_raw, total)
            if wanted is not None and not wanted:
                return "", "That page range is outside the document."
            chunks = []
            for number, page in enumerate(pdf.pages, start=1):
                if wanted is not None and number not in wanted:
                    continue
                try:
                    chunks.append(page.extract_text() or "")
                except Exception:
                    chunks.append("")
        text = "\n".join(chunks)
        if text.strip():
            return text, ""
    except ImportError:
        pass
    except Exception as exc:
        return "", f"I could not read that PDF: {type(exc).__name__}."

    # Tier 3: PyPDF2 fallback
    try:
        import PyPDF2
        with open(path, "rb") as handle:          # read-only, never written back
            reader = PyPDF2.PdfReader(handle)
            total = len(reader.pages)
            wanted = _page_set(pages_raw, total)
            if wanted is not None and not wanted:
                return "", "That page range is outside the document."
            chunks = []
            for number, page in enumerate(reader.pages, start=1):
                if wanted is not None and number not in wanted:
                    continue
                try:
                    chunks.append(page.extract_text() or "")
                except Exception:
                    chunks.append("")
        return "\n".join(chunks), ""
    except ImportError:
        return "", "No PDF reader is installed, so I cannot open that file."
    except Exception as exc:
        return "", f"I could not read that PDF: {type(exc).__name__}."


def _extract_docx(path: Path) -> tuple[str, str]:
    try:
        import docx
    except ImportError:
        return "", "The Word reader is not installed."
    try:
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts), ""
    except Exception as exc:
        return "", f"I could not read that Word file: {type(exc).__name__}."


def _extract_xlsx(path: Path) -> tuple[str, str]:
    try:
        import openpyxl
    except ImportError:
        return "", "The spreadsheet reader is not installed."
    try:
        # read_only + data_only: no formulas are evaluated and nothing is written
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        try:
            for sheet in workbook.worksheets:
                parts.append(f"### Sheet: {sheet.title}")
                for count, row in enumerate(sheet.iter_rows(values_only=True)):
                    if count >= _MAX_SHEET_ROWS:
                        parts.append("… (sheet truncated)")
                        break
                    cells = ["" if v is None else str(v) for v in row]
                    if any(c.strip() for c in cells):
                        parts.append(" | ".join(cells))
        finally:
            workbook.close()
        return "\n".join(parts), ""
    except Exception as exc:
        return "", f"I could not read that spreadsheet: {type(exc).__name__}."


def _extract_pptx(path: Path) -> tuple[str, str]:
    try:
        from pptx import Presentation
    except ImportError:
        return "", "The PowerPoint reader is not installed."
    try:
        deck = Presentation(str(path))
        parts = []
        for number, slide in enumerate(deck.slides, start=1):
            parts.append(f"### Slide {number}")
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    for paragraph in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in paragraph.runs).strip()
                        if line:
                            parts.append(line)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        if any(cells):
                            parts.append(" | ".join(cells))
        return "\n".join(parts), ""
    except Exception as exc:
        return "", f"I could not read that presentation: {type(exc).__name__}."


def _extract_plain(path: Path) -> tuple[str, str]:
    try:
        return path.read_text(encoding="utf-8", errors="replace"), ""
    except Exception as exc:
        return "", f"I could not read that file: {type(exc).__name__}."


def _extract(path: Path, pages_raw):
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path, pages_raw)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in (".xlsx", ".xlsm"):
        return _extract_xlsx(path)
    if suffix == ".pptx":
        return _extract_pptx(path)
    if suffix in _PLAIN_SUFFIXES:
        return _extract_plain(path)
    return "", (f"I can read PDF, Word, Excel, PowerPoint and plain-text files — "
                f"'{suffix or path.name}' is not one of them.")


def document_qa_action(parameters: dict, player=None) -> str:
    """
    Answer a question about a document. Read-only.

    parameters:
        file_path : document to read; falls back to the file currently open in the UI
        question  : what the user wants to know (required)
        pages     : optional page restriction for PDFs, e.g. '3', '1-5', '2,4,6-8'
    """
    p = parameters if isinstance(parameters, dict) else {}

    question = str(p.get("question") or "").strip()
    if not question:
        return "What would you like me to look for in the document?"

    raw_path = p.get("file_path") or p.get("path") or p.get("file")
    if not str(raw_path or "").strip() and player is not None:
        # Same convenience main.py already gives file_processor: no path spoken,
        # so use the document the user has open.
        raw_path = getattr(player, "current_file", None)

    path = _safe_file(raw_path)
    if path is None:
        return ("I need a document to read, and it has to be inside your own "
                "folders. Which file do you mean?")
    if not path.exists() or not path.is_file():
        return f"I cannot find that document: {path.name}"
    if path.suffix.lower() not in _KNOWN_SUFFIXES:
        return (f"I can read PDF, Word, Excel, PowerPoint and plain-text files — "
                f"'{path.suffix or path.name}' is not one of them.")
    try:
        if path.stat().st_size > _MAX_FILE_BYTES:
            return "That document is too large for me to read in one go."
    except OSError:
        return "I cannot access that document."

    text, error = _extract(path, p.get("pages"))
    if error:
        return error
    if not text.strip():
        return (f"'{path.name}' has no extractable text — it may be a scan. "
                "Try screen_process on the open document instead.")

    body = text[:_MAX_CHARS]
    truncated = len(text) > _MAX_CHARS
    prompt = (
        f"Answer the question using ONLY the document excerpt below. If the "
        f"answer is not in the excerpt, say so plainly instead of guessing, and "
        f"never invent figures, dates, names or clauses. Quote the document "
        f"briefly where it helps. Reply in the language of the question.\n\n"
        f"QUESTION: {question}\n\n"
        f"DOCUMENT '{path.name}'"
        f"{' (excerpt, the document was cut short)' if truncated else ''}:\n"
        f"--- DOCUMENT START ---\n{body}\n--- DOCUMENT END ---"
    )

    try:
        from core import gemini
    except Exception as exc:
        return f"The document reader is unavailable: {type(exc).__name__}."

    answer = gemini.text(prompt, tier=gemini.SMART, timeout_ms=_ANSWER_TIMEOUT_MS,
                         default="").strip()
    if not answer:
        return "I read the document but the model did not answer — please try again."

    if player is not None:
        try:
            player.show_content(f"{path.name[:30]} — Q&A", f"{question}\n\n{answer}")
        except Exception:
            pass
        try:
            player.write_log(f"DOC-QA: {path.name}, {len(body)} chars read"
                             + (" (truncated)" if truncated else ""))
        except Exception:
            pass

    return f"From {path.name}: {answer}"


TOOL = {
    "name": "document_qa",
    "description": (
        "Answers a question about an existing PDF, Word, Excel, PowerPoint or "
        "plain-text document. Strictly read-only: it never edits, converts or "
        "creates files. Use it when the user asks what a document says, wants a "
        "clause, figure, table cell, slide or section found, or asks a question "
        "about the file currently open. For converting, resizing, splitting or "
        "otherwise changing a file use file_processor; for summarizing a web page "
        "or the clipboard use summarize."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "file_path": {
                "type": "STRING",
                "description": "Path of the document to read. If omitted, the file "
                               "currently open in the app is used.",
            },
            "question": {
                "type": "STRING",
                "description": "What the user wants to know about the document.",
            },
            "pages": {
                "type": "STRING",
                "description": "Optional page restriction for PDFs, e.g. '3', "
                               "'1-5' or '2,4,6-8'. Ignored for other formats.",
            },
        },
        "required": ["question"],
    },
    "handler": document_qa_action,
    "behavior": "NON_BLOCKING",
    "scheduling": "WHEN_IDLE",
}
